"""Build the DHVAJ Client Master workbook — the template the portal's
"Export client master" fills and "Import client master" reads.

    python3 -m venv /tmp/cm && /tmp/cm/bin/pip install openpyxl
    /tmp/cm/bin/python scripts/client-master/build_template.py \
        apps/api/src/modules/credentials/client-master-template.xlsx

The headings on row 4 of 'Clients' and row 3 of 'Logins' are the contract with
apps/api/src/modules/credentials/client-master.ts (CLIENT_MASTER_COLUMNS in
@hsdg/contracts): change one, change both. Version: client-master v2 (v2 added the
login columns AF:AQ; a v1 file still imports).
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
FIRST, LAST = 6, 1005  # data rows on Clients (row 5 = example)
FONT = "Arial"

# ── Vocabulary (mirrors packages/contracts) ──────────────────────────────
ENTITY_TYPES = [  # name, slug, category
    ("Private Limited Company", "private_limited", "company"),
    ("Public Limited Company", "public_limited", "company"),
    ("One Person Company", "opc", "company"),
    ("Limited Liability Partnership", "llp", "llp"),
    ("Partnership Firm", "partnership", "partnership"),
    ("Proprietorship", "proprietorship", "proprietorship"),
    ("Individual", "individual", "individual"),
    ("Hindu Undivided Family", "huf", "huf"),
    ("Trust", "trust", "trust"),
    ("Society", "society", "society"),
    ("Association of Persons", "aop", "aop"),
]
PAN_GUESS = [  # 4th PAN letter → type (entityTypeFromPan)
    ("C", "Private Limited Company"),
    ("P", "Proprietorship"),
    ("F", "Partnership Firm"),
    ("H", "Hindu Undivided Family"),
    ("A", "Association of Persons"),
    ("B", "Association of Persons"),
    ("T", "Trust"),
]
LEGAL_STATUSES = [
    ("Active", "active"), ("Under incorporation", "under_incorporation"),
    ("Dormant", "dormant"), ("Struck off", "struck_off"),
    ("Under liquidation", "under_liquidation"), ("CIRP", "cirp"),
    ("Closed", "closed"), ("Other", "other"),
]
STATES = {
    "01": "Jammu and Kashmir", "02": "Himachal Pradesh", "03": "Punjab", "04": "Chandigarh",
    "05": "Uttarakhand", "06": "Haryana", "07": "Delhi", "08": "Rajasthan", "09": "Uttar Pradesh",
    "10": "Bihar", "11": "Sikkim", "12": "Arunachal Pradesh", "13": "Nagaland", "14": "Manipur",
    "15": "Mizoram", "16": "Tripura", "17": "Meghalaya", "18": "Assam", "19": "West Bengal",
    "20": "Jharkhand", "21": "Odisha", "22": "Chhattisgarh", "23": "Madhya Pradesh", "24": "Gujarat",
    "26": "Dadra and Nagar Haveli and Daman and Diu", "27": "Maharashtra", "29": "Karnataka",
    "30": "Goa", "31": "Lakshadweep", "32": "Kerala", "33": "Tamil Nadu", "34": "Puducherry",
    "35": "Andaman and Nicobar Islands", "36": "Telangana", "37": "Andhra Pradesh", "38": "Ladakh",
    "97": "Other Territory",
}
FYS = ["2024-25", "2025-26", "2026-27", "2027-28"]
PORTALS = ["GST", "Income-tax (PAN)", "Income-tax (TAN)", "TRACES deductor",
           "TRACES taxpayer", "MCA", "E-way bill", "E-invoice"]
GST_CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin_check(first14: str) -> str:
    s = 0
    for i, ch in enumerate(first14):
        p = GST_CHARS.index(ch) * (1 if i % 2 == 0 else 2)
        s += p // 36 + p % 36
    return GST_CHARS[(36 - s % 36) % 36]


EX_PAN = "ABCPS1234K"
EX_GSTIN = "07" + EX_PAN + "1Z"
EX_GSTIN += gstin_check(EX_GSTIN)

# ── Styles ────────────────────────────────────────────────────────────────
def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)

INPUT = fill("FFF8D6")      # yellow-ish: you fill
AUTO = fill("ECEFF3")       # grey: fills itself
EXAMPLE = fill("E3F2FD")
TITLE = fill("1F3A5F")
GROUPS = {
    "Client": fill("1F3A5F"), "Registrations": fill("2E5E4E"), "Contact": fill("5B4B8A"),
    "Address": fill("7A4E2D"), "Engagement": fill("8A2D4E"), "Logins": fill("3D5A80"),
    "Checks (fill themselves)": fill("4A4A4A"),
}
thin = Side(style="thin", color="C9CED6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)

wb = Workbook()

# ── Lists ─────────────────────────────────────────────────────────────────
ls = wb.active
ls.title = "Lists"
lists_cols = {}

def put_list(col, header, rows):
    ls.cell(row=1, column=col, value=header).font = f(bold=True, color="FFFFFF")
    ls.cell(row=1, column=col).fill = TITLE
    for i, r in enumerate(rows, start=2):
        vals = r if isinstance(r, tuple) else (r,)
        for j, v in enumerate(vals):
            c = ls.cell(row=i, column=col + j, value=v)
            c.font = f()
            if isinstance(v, str) and v.isdigit():
                c.number_format = "@"
    return 2, 1 + len(rows)

r0, r1 = put_list(1, "Entity type", ENTITY_TYPES)
ls["B1"], ls["C1"] = "Portal slug", "Category"
for c in ("B1", "C1"):
    ls[c].font, ls[c].fill = f(bold=True, color="FFFFFF"), TITLE
TYPE_RANGE = f"Lists!$A$2:$A${r1}"
TYPE_TABLE = f"Lists!$A$2:$C${r1}"

r0, r1 = put_list(5, "PAN 4th letter", PAN_GUESS)
ls["F1"] = "Guessed type"
ls["F1"].font, ls["F1"].fill = f(bold=True, color="FFFFFF"), TITLE
GUESS_TABLE = f"Lists!$E$2:$F${r1}"

r0, r1 = put_list(8, "Legal status", LEGAL_STATUSES)
ls["I1"] = "Portal value"
ls["I1"].font, ls["I1"].fill = f(bold=True, color="FFFFFF"), TITLE
LEGAL_RANGE = f"Lists!$H$2:$H${r1}"

r0, r1 = put_list(11, "GST state code", list(STATES.items()))
ls["L1"] = "State"
ls["L1"].font, ls["L1"].fill = f(bold=True, color="FFFFFF"), TITLE
STATE_CODES, STATE_NAMES = f"Lists!$K$2:$K${r1}", f"Lists!$L$2:$L${r1}"
STATE_NAME_LIST = STATE_NAMES

_, r1 = put_list(14, "Financial year", FYS)
FY_RANGE = f"Lists!$N$2:$N${r1}"
_, r1 = put_list(16, "Yes / No", ["Yes", "No"])
YN_RANGE = f"Lists!$P$2:$P${r1}"
_, r1 = put_list(18, "GST frequency", ["Monthly", "Quarterly"])
FREQ_RANGE = f"Lists!$R$2:$R${r1}"
_, r1 = put_list(20, "Login portal", PORTALS)
PORTAL_RANGE = f"Lists!$T$2:$T${r1}"
for col, w in {"A": 30, "B": 16, "C": 14, "E": 15, "F": 24, "H": 20, "I": 20,
               "K": 15, "L": 38, "N": 15, "P": 10, "R": 15, "T": 20}.items():
    ls.column_dimensions[col].width = w
ls.sheet_properties.tabColor = "9AA5B1"

# ── Clients ───────────────────────────────────────────────────────────────
LOGIN_PAIRS = [
    ("GST user ID", "GST password", "The main GSTIN's GST portal login. Other GSTINs' logins: 'Logins' sheet."),
    ("Income-tax user ID (PAN)", "Income-tax password (PAN)", "Income-tax e-filing login of the PAN (user ID is usually the PAN)."),
    ("Income-tax user ID (TAN)", "Income-tax password (TAN)", "Income-tax e-filing login of the main TAN."),
    ("TRACES user ID (TAN)", "TRACES password", "TRACES deductor login of the main TAN."),
    ("MCA user ID", "MCA password", "MCA V3 login used for the company / LLP."),
    ("E-way bill user ID", "E-way bill password", "E-way bill portal login of the main GSTIN."),
]
LOGIN_HEADERS = {h for u, pw, _ in LOGIN_PAIRS for h in (u, pw)}
cs = wb.create_sheet("Clients", 0)
# (group, header, width, kind, note)   kind: in | auto | hide
COLS = [
    ("Client", "Client code", 12, "in", "Filled by Export for clients already in the portal — leave it as it is. Blank = a NEW client."),
    ("Client", "Legal name *", 34, "in", "Required. As on the PAN / GST certificate."),
    ("Client", "Trade name", 22, "in", "Brand / trade name, if different."),
    ("Client", "Entity type", 24, "in", "Pick from the list. Blank = guessed from the PAN (see 'Entity type to use')."),
    ("Registrations", "PAN", 13, "in", "10 characters. Blank = taken from the GSTIN."),
    ("Registrations", "GSTIN (main)", 18, "in", "15 characters. PAN, state and type are worked out from it."),
    ("Registrations", "Other GSTINs", 22, "in", "Other states' GSTINs of the SAME client, separated by commas."),
    ("Registrations", "TAN", 13, "in", "Needed for a TDS/TCS engagement."),
    ("Registrations", "CIN / LLPIN", 22, "in", "Needed for an ROC/MCA engagement (company / LLP)."),
    ("Registrations", "Date of incorporation / birth", 14, "in", "dd-mm-yyyy"),
    ("Registrations", "Legal status", 16, "in", "Blank = Active."),
    ("Registrations", "Home office code", 11, "in", "Blank = the firm's main office."),
    ("Contact", "Contact person", 20, "in", "Primary contact."),
    ("Contact", "Designation", 14, "in", ""),
    ("Contact", "Email", 24, "in", ""),
    ("Contact", "Mobile", 14, "in", ""),
    ("Address", "Address line 1", 26, "in", "Registered / principal place of business."),
    ("Address", "Address line 2", 20, "in", ""),
    ("Address", "City", 14, "in", ""),
    ("Address", "State", 18, "in", "Blank = the GSTIN's state."),
    ("Address", "Pincode", 9, "in", "6 digits."),
    ("Engagement", "Financial year", 11, "in", "Year the engagement starts, e.g. 2026-27."),
    ("Engagement", "GST", 7, "in", "Yes = GST compliance engagement (returns calendar)."),
    ("Engagement", "GST return frequency", 12, "in", "Monthly or Quarterly (QRMP)."),
    ("Engagement", "TDS / TCS", 8, "in", ""),
    ("Engagement", "Income Tax", 8, "in", ""),
    ("Engagement", "ROC / MCA", 8, "in", "Companies and LLPs only."),
    ("Engagement", "Bookkeeping", 10, "in", ""),
    ("Engagement", "MIS", 7, "in", ""),
    ("Engagement", "Engagement partner (email)", 24, "in", "Portal login email. Blank = you."),
    ("Engagement", "Lead / in-charge (email)", 24, "in", "Portal login email of the staff member running it."),
    # (user ID header, password header, note) — the client's main logins (CLIENT_MASTER_LOGIN_CELLS)
    *[c for u, pw, note in LOGIN_PAIRS
      for c in (("Logins", u, 18, "in", note), ("Logins", pw, 16, "in", "Encrypted when imported. "
                "Export shows '(saved in portal)' unless you choose Export with passwords."))],
    ("Checks (fill themselves)", "Row type", 13, "auto", ""),
    ("Checks (fill themselves)", "PAN to use", 13, "auto", "Your PAN, or the one inside the GSTIN."),
    ("Checks (fill themselves)", "Entity type to use", 26, "auto", ""),
    ("Checks (fill themselves)", "GST state", 18, "auto", "From the first two digits of the GSTIN."),
    ("Checks (fill themselves)", "GSTIN check", 16, "auto", "Checks the format, state code and the check digit."),
    ("Checks (fill themselves)", "PAN check", 13, "auto", ""),
    ("Checks (fill themselves)", "TAN check", 13, "auto", ""),
    ("Checks (fill themselves)", "Ready for engagement?", 48, "auto", "What is still missing before the import can create the client and its engagement."),
    # hidden helpers
    ("", "_gstin", 4, "hide", ""), ("", "_pan", 4, "hide", ""), ("", "_type", 4, "hide", ""),
    ("", "_cat", 4, "hide", ""), ("", "_streams", 4, "hide", ""), ("", "_issues", 4, "hide", ""),
]
L = {h: get_column_letter(i) for i, (_, h, *_rest) in enumerate(COLS, start=1)}
NCOL = len(COLS)
LAST_VISIBLE = max(i for i, c in enumerate(COLS, start=1) if c[3] != "hide")

cs["A1"] = "DHVAJ Client Master — fill once, import into the portal, export any time"
cs["A1"].font = f(bold=True, color="FFFFFF", size=14)
cs["A1"].fill = TITLE
cs.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_VISIBLE)
cs.row_dimensions[1].height = 26
cs["A2"] = ("Yellow = you fill   ·   Grey = fills itself (do not type there)   ·   Blue row 5 = example, "
            "ignored by the import   ·   * = required   ·   See 'Read Me' for import / export steps")
cs["A2"].font = f(italic=True, color="333333")
cs.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_VISIBLE)

# group band (row 3) and headers (row 4)
start = 1
for i in range(1, NCOL + 2):
    g = COLS[i - 1][0] if i <= NCOL else None
    if i > NCOL or g != COLS[start - 1][0]:
        grp = COLS[start - 1][0]
        if grp:
            cs.cell(row=3, column=start, value=grp)
            for c in range(start, i):
                cs.cell(row=3, column=c).fill = GROUPS[grp]
            cs.merge_cells(start_row=3, start_column=start, end_row=3, end_column=i - 1)
            cs.cell(row=3, column=start).font = f(bold=True, color="FFFFFF")
            cs.cell(row=3, column=start).alignment = CENTER
        start = i
for i, (grp, h, w, kind, note) in enumerate(COLS, start=1):
    c = cs.cell(row=4, column=i, value=h)
    c.font = f(bold=True, color="FFFFFF" if grp else "999999")
    c.fill = GROUPS.get(grp, fill("FFFFFF"))
    c.alignment = CENTER
    c.border = BORDER
    if note:
        c.comment = Comment(note, "DHVAJ")
    col = get_column_letter(i)
    cs.column_dimensions[col].width = w
    if kind == "hide":
        cs.column_dimensions[col].hidden = True
cs.row_dimensions[4].height = 42


def formulas(r):
    c = {h: f"{L[h]}{r}" for h in L}
    g, p, t, cat, ns, iss = c["_gstin"], c["_pan"], c["_type"], c["_cat"], c["_streams"], c["_issues"]
    filled = f'COUNTA({L["Client code"]}{r}:{L["Lead / in-charge (email)"]}{r})>0'
    yes = lambda h: f'{c[h]}="Yes"'
    pos14 = "{1,2,3,4,5,6,7,8,9,10,11,12,13,14}"
    wts14 = "{1,2,1,2,1,2,1,2,1,2,1,2,1,2}"
    val14 = f'(FIND(MID({g},{pos14},1),"{GST_CHARS}")-1)*{wts14}'
    checksum = f'MID("{GST_CHARS}",MOD(36-MOD(SUMPRODUCT(INT({val14}/36)+MOD({val14},36)),36),36)+1,1)'
    letters = '"ABCDEFGHIJKLMNOPQRSTUVWXYZ"'
    def pan_chars(x):  # x: a 10-char text expression → letters 1-5,10 and digits 6-9
        return (f'SUMPRODUCT(--ISERROR(FIND(MID({x},{{1,2,3,4,5,10}},1),{letters})))=0,'
                f'SUMPRODUCT(--ISERROR(FIND(MID({x},{{6,7,8,9}},1),"0123456789")))=0')
    def pan_shape(x):
        return f'AND(LEN({x})=10,{pan_chars(x)})'
    gst_shape = (f'AND(LEN({g})=15,ISNUMBER(MATCH(LEFT({g},2),{STATE_CODES},0)),'
                 f'{pan_chars(f"MID({g},3,10)")},'
                 f'ISNUMBER(FIND(MID({g},13,1),"123456789{letters[1:-1]}")),MID({g},14,1)="Z")')
    tan_shape = (f'AND(LEN(UPPER(TRIM({c["TAN"]})))=10,'
                 f'SUMPRODUCT(--ISERROR(FIND(MID(UPPER(TRIM({c["TAN"]})),{{1,2,3,4,10}},1),{letters})))=0,'
                 f'SUMPRODUCT(--ISERROR(FIND(MID(UPPER(TRIM({c["TAN"]})),{{5,6,7,8,9}},1),"0123456789")))=0)')
    pan_use = c["PAN to use"]
    out = {
        "_gstin": f'=UPPER(SUBSTITUTE(TRIM({c["GSTIN (main)"]})," ",""))',
        "_pan": f'=UPPER(SUBSTITUTE(TRIM({c["PAN"]})," ",""))',
        "_type": (f'=IF({c["Entity type"]}<>"",{c["Entity type"]},'
                  f'IFERROR(VLOOKUP(MID({pan_use},4,1),{GUESS_TABLE},2,FALSE),""))'),
        "_cat": f'=IFERROR(VLOOKUP({t},{TYPE_TABLE},3,FALSE),"")',
        "_streams": f'=COUNTIF({L["GST"]}{r}:{L["MIS"]}{r},"Yes")',
        "Row type": f'=IF(NOT({filled}),"",IF({c["Client code"]}="","New client","Update"))',
        "PAN to use": f'=IF({p}<>"",{p},IF(LEN({g})=15,MID({g},3,10),""))',
        "Entity type to use": (f'=IF({t}="","",{t}&IF({c["Entity type"]}="","  (guessed from PAN)",""))'),
        "GST state": f'=IF({g}="","",IFERROR(INDEX({STATE_NAMES},MATCH(LEFT({g},2),{STATE_CODES},0)),"Unknown code"))',
        "GSTIN check": (f'=IF({g}="","",IFERROR(IF(NOT({gst_shape}),"Wrong format",'
                        f'IF({checksum}=MID({g},15,1),"OK","Typo? check digit wrong")),"Wrong format"))'),
        "PAN check": f'=IF({p}="","",IF({pan_shape(p)},"OK","Wrong format"))',
        "TAN check": f'=IF(TRIM({c["TAN"]})="","",IF({tan_shape},"OK","Wrong format"))',
        "_issues": (
            '=' + '&'.join([
                f'IF(TRIM({c["Legal name *"]})="","Legal name missing; ","")',
                f'IF(AND({c["GSTIN check"]}<>"",{c["GSTIN check"]}<>"OK"),"GSTIN: "&{c["GSTIN check"]}&"; ","")',
                f'IF(AND({c["PAN check"]}<>"",{c["PAN check"]}<>"OK"),"PAN wrong format; ","")',
                f'IF(AND({p}<>"",LEN({g})=15,MID({g},3,10)<>{p}),"PAN does not match GSTIN; ","")',
                f'IF(AND({c["TAN check"]}<>"",{c["TAN check"]}<>"OK"),"TAN wrong format; ","")',
                f'IF({pan_use}="","PAN missing; ","")',
                f'IF({t}="","Entity type missing; ","")',
                f'IF(AND({g}<>"",COUNTIF(${L["_gstin"]}${FIRST}:${L["_gstin"]}${LAST},{g})>1),"GSTIN repeated in this sheet; ","")',
                f'IF(AND({pan_use}<>"",COUNTIF(${L["PAN to use"]}${FIRST}:${L["PAN to use"]}${LAST},{pan_use})>1),"PAN repeated in this sheet (same client? put extra GSTINs in Other GSTINs); ","")',
                f'IF(AND({ns}>0,{c["Financial year"]}=""),"Financial year missing; ","")',
                f'IF(AND({ns}=0,{c["Financial year"]}<>""),"Financial year given but no service is Yes; ","")',
                f'IF(AND({yes("GST")},{g}=""),"GST needs a GSTIN; ","")',
                f'IF(AND({yes("GST")},{c["GST return frequency"]}=""),"GST return frequency missing; ","")',
                f'IF(AND({yes("TDS / TCS")},TRIM({c["TAN"]})=""),"TDS needs a TAN; ","")',
                f'IF(AND({yes("ROC / MCA")},{cat}<>"",{cat}<>"company",{cat}<>"llp"),"ROC is only for a company / LLP; ","")',
                f'IF(AND({yes("ROC / MCA")},OR({cat}="company",{cat}="llp"),TRIM({c["CIN / LLPIN"]})=""),"ROC needs the CIN / LLPIN; ","")',
                *[f'IF((TRIM({c[u]})="")<>(TRIM({c[pw]})=""),"{u.split(" user ID")[0]} login needs both user ID and password; ","")'
                  for u, pw, _ in LOGIN_PAIRS],
            ])
        ),
        "Ready for engagement?": (f'=IF(NOT({filled}),"",IF({iss}<>"","✖ "&LEFT({iss},LEN({iss})-2),'
                                  f'IF({ns}>0,"✔ Ready — client + "&{ns}&" service(s)","✔ Ready — client only")))'),
    }
    return out


def write_row(r, values=None, example=False):
    fm = formulas(r)
    for i, (grp, h, w, kind, note) in enumerate(COLS, start=1):
        cell = cs.cell(row=r, column=i)
        if kind == "in":
            if values and h in values:
                cell.value = values[h]
            cell.fill = EXAMPLE if example else INPUT
            cell.font = f(italic=example, color="1F3A5F" if example else "000000")
        else:
            cell.value = fm.get(h)
            cell.fill = AUTO if kind == "auto" else fill("FFFFFF")
            cell.font = f(color="333333")
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(h == "Ready for engagement?"))
        if h in ("PAN", "GSTIN (main)", "TAN", "Pincode", "Mobile", "Client code", "CIN / LLPIN",
                 "Other GSTINs", "Home office code") or h in LOGIN_HEADERS:
            cell.number_format = "@"
        if h == "Date of incorporation / birth":
            cell.number_format = "DD-MM-YYYY"


write_row(5, {
    "Client code": "EXAMPLE", "Legal name *": "Sharma Traders", "Trade name": "Sharma Electricals",
    "Entity type": "", "PAN": "", "GSTIN (main)": EX_GSTIN, "Other GSTINs": "", "TAN": "DELS12345A",
    "CIN / LLPIN": "", "Date of incorporation / birth": date(2015, 4, 10), "Legal status": "Active",
    "Home office code": "", "Contact person": "Rakesh Sharma", "Designation": "Proprietor",
    "Email": "accounts@sharmatraders.example", "Mobile": "9810000000", "Address line 1": "12, Main Market",
    "Address line 2": "Karol Bagh", "City": "New Delhi", "State": "", "Pincode": "110005",
    "Financial year": "2026-27", "GST": "Yes", "GST return frequency": "Monthly", "TDS / TCS": "Yes",
    "Income Tax": "Yes", "ROC / MCA": "No", "Bookkeeping": "No", "MIS": "No",
    "Engagement partner (email)": "", "Lead / in-charge (email)": "",
    "GST user ID": "sharma_gst", "GST password": "••••••••",
    "Income-tax user ID (PAN)": EX_PAN, "Income-tax password (PAN)": "••••••••",
}, example=True)
cs.cell(row=5, column=1).comment = Comment(
    "Example row — the import skips any row whose Client code is EXAMPLE. Overwrite or delete it.", "DHVAJ")
for r in range(FIRST, LAST + 1):
    write_row(r)

# Row type for the example row
cs[f'{L["Row type"]}5'] = '=IF(A5="EXAMPLE","Example (skipped)","")'

# validations
def dv_list(src, col, allow_blank=True, error=None):
    dv = DataValidation(type="list", formula1=f"={src}", allow_blank=allow_blank, showDropDown=False)
    dv.error = error or "Pick a value from the list."
    dv.errorTitle = "Not in the list"
    dv.showErrorMessage = True
    cs.add_data_validation(dv)
    dv.add(f"{L[col]}5:{L[col]}{LAST}")

dv_list(TYPE_RANGE, "Entity type")
dv_list(LEGAL_RANGE, "Legal status")
dv_list(STATE_NAME_LIST, "State")
dv_list(FY_RANGE, "Financial year")
dv_list(FREQ_RANGE, "GST return frequency")
for h in ("GST", "TDS / TCS", "Income Tax", "ROC / MCA", "Bookkeeping", "MIS"):
    dv_list(YN_RANGE, h)
dv = DataValidation(type="date", operator="between", formula1="DATE(1900,1,1)", formula2="DATE(2100,12,31)",
                    allow_blank=True)
dv.error, dv.errorTitle, dv.showErrorMessage = "Type a date, e.g. 10-04-2015.", "Not a date", True
cs.add_data_validation(dv)
dv.add(f'{L["Date of incorporation / birth"]}5:{L["Date of incorporation / birth"]}{LAST}')
for h, n in (("PAN", 10), ("GSTIN (main)", 15), ("TAN", 10), ("Pincode", 6)):
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(n + 2), allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = f"{h} is {n} characters.", "Too long", True
    cs.add_data_validation(dv)
    dv.add(f"{L[h]}5:{L[h]}{LAST}")

# conditional formatting: checks
rng_ready = f'{L["Ready for engagement?"]}5:{L["Ready for engagement?"]}{LAST}'
top = f'{L["Ready for engagement?"]}5'
cs.conditional_formatting.add(rng_ready, FormulaRule(formula=[f'LEFT({top},1)="✔"'],
                              fill=fill("D7F0DD"), font=Font(name=FONT, color="1B5E20", bold=True)))
cs.conditional_formatting.add(rng_ready, FormulaRule(formula=[f'LEFT({top},1)="✖"'],
                              fill=fill("FBE0E0"), font=Font(name=FONT, color="B71C1C")))
for h in ("GSTIN check", "PAN check", "TAN check"):
    rng = f"{L[h]}5:{L[h]}{LAST}"
    t = f"{L[h]}5"
    cs.conditional_formatting.add(rng, FormulaRule(formula=[f'{t}="OK"'], font=Font(name=FONT, color="1B5E20", bold=True)))
    cs.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({t}<>"",{t}<>"OK")'],
                                  fill=fill("FBE0E0"), font=Font(name=FONT, color="B71C1C", bold=True)))
et = f'{L["Entity type to use"]}5'
cs.conditional_formatting.add(f'{L["Entity type to use"]}5:{L["Entity type to use"]}{LAST}',
                              FormulaRule(formula=[f'ISNUMBER(SEARCH("guessed",{et}))'],
                                          font=Font(name=FONT, color="8A5A00", italic=True)))

cs.freeze_panes = "C5"
cs.auto_filter.ref = f"A4:{get_column_letter(LAST_VISIBLE)}{LAST}"
cs.sheet_properties.tabColor = "1F3A5F"

# ── Logins ────────────────────────────────────────────────────────────────
lg = wb.create_sheet("Logins", 1)
lg["A1"] = "More portal logins (optional) — other GSTINs, TRACES taxpayer, E-invoice; one row per login"
lg["A1"].font, lg["A1"].fill = f(bold=True, color="FFFFFF", size=13), TITLE
lg.merge_cells("A1:G1")
lg["A2"] = ("A client's main logins go on its own row on 'Clients' (Logins columns). Passwords are sealed "
            "(encrypted) when imported. Export masks them unless you choose Export with passwords — such a file "
            "holds real passwords: share it only with people who may see them, and delete it after use.")
lg["A2"].font = f(italic=True, color="8A2D2D")
lg["A2"].alignment = WRAP
lg.merge_cells("A2:G2")
lg.row_dimensions[2].height = 30
LH = [("Client code", 12), ("Legal name (as on Clients)", 32), ("Portal", 18),
      ("For GSTIN / TAN", 18), ("User ID", 22), ("Password", 18), ("Check", 34)]
for i, (h, w) in enumerate(LH, start=1):
    c = lg.cell(row=3, column=i, value=h)
    c.font, c.fill, c.alignment, c.border = f(bold=True, color="FFFFFF"), GROUPS["Client"] if i < 7 else GROUPS["Checks (fill themselves)"], CENTER, BORDER
    lg.column_dimensions[get_column_letter(i)].width = w
lg["C3"].comment = Comment("GST login goes with a GSTIN; Income-tax (PAN) with the client's PAN; "
                           "Income-tax (TAN) and TRACES deductor with the TAN.", "DHVAJ")
lg["D3"].comment = Comment("Which GSTIN (for a GST login) or TAN (for TAN / TRACES deductor). "
                           "Blank = the client's main GSTIN / TAN.", "DHVAJ")
NAME_COL, CODE_COL = f"Clients!${L['Legal name *']}$5:${L['Legal name *']}${LAST}", f"Clients!${L['Client code']}$5:${L['Client code']}${LAST}"
GST_ALL = f"Clients!${L['_gstin']}$5:${L['_gstin']}${LAST}"
OTHER_ALL = f"Clients!${L['Other GSTINs']}$5:${L['Other GSTINs']}${LAST}"
LOG_LAST = 1503
for r in range(4, LOG_LAST + 1):
    ex = r == 4
    for i in range(1, 8):
        c = lg.cell(row=r, column=i)
        c.border = BORDER
        c.fill = AUTO if i == 7 else (EXAMPLE if ex else INPUT)
        c.font = f(italic=ex, color="1F3A5F" if ex else "000000")
        c.number_format = "@" if i in (1, 4, 5, 6) else "General"
    lg.cell(row=r, column=7).value = (
        f'=IF(COUNTA(A{r}:F{r})=0,"",IF(A{r}="EXAMPLE","Example (skipped)",'
        f'IF(OR(C{r}="",E{r}="",F{r}=""),"✖ Portal, User ID and Password are all needed",'
        f'IF(AND(A{r}="",COUNTIF({NAME_COL},B{r})=0),"Not on Clients — will match a portal client by name",'
        f'IF(AND(C{r}="GST",D{r}<>"",COUNTIF({GST_ALL},UPPER(D{r}))=0,COUNTIF({OTHER_ALL},"*"&UPPER(D{r})&"*")=0),'
        f'"GSTIN not on Clients — will be checked in the portal","✔ OK")))))'
    )
    lg.cell(row=r, column=7).font = f(color="333333")
for i, v in enumerate(["EXAMPLE", "Sharma Traders", "GST", EX_GSTIN, "sharma_gst", "••••••••"], start=1):
    lg.cell(row=4, column=i, value=v)
dv = DataValidation(type="list", formula1=f"={PORTAL_RANGE}", allow_blank=True)
dv.error, dv.errorTitle, dv.showErrorMessage = "Pick a portal from the list.", "Not in the list", True
lg.add_data_validation(dv)
dv.add(f"C4:C{LOG_LAST}")
lg.conditional_formatting.add(f"G4:G{LOG_LAST}", FormulaRule(formula=['LEFT(G4,1)="✔"'], font=Font(name=FONT, color="1B5E20", bold=True)))
lg.conditional_formatting.add(f"G4:G{LOG_LAST}", FormulaRule(formula=['LEFT(G4,1)="✖"'], fill=fill("FBE0E0"), font=Font(name=FONT, color="B71C1C")))
lg.freeze_panes = "C4"
lg.sheet_properties.tabColor = "5B4B8A"

# ── Read Me ───────────────────────────────────────────────────────────────
rm = wb.create_sheet("Read Me", 0)
rm.column_dimensions["A"].width = 4
rm.column_dimensions["B"].width = 30
rm.column_dimensions["C"].width = 95
rows = [
    ("h1", "DHVAJ Client Master", ""),
    ("p", "", "One sheet for every client's master details, so you stop collecting GSTIN / PAN / TAN from "
              "different sites. Fill a row, and the grey columns check it and tell you what is still missing "
              "to open the client's engagements."),
    ("", "", ""),
    ("h2", "Colours", ""),
    ("in", "Yellow cell", "You fill it."),
    ("auto", "Grey cell", "Fills itself from what you typed — do not type there."),
    ("ex", "Blue row", "Example (Client code EXAMPLE). The import skips it."),
    ("", "", ""),
    ("h2", "Export (portal → Excel)", ""),
    ("p", "1.", "Clients → Export client master. You get this same file with every client "
                "already in the portal, one per row, with its Client code."),
    ("p", "2.", "Use the 'Ready for engagement?' column (or filter it on ✖) to see which clients still have "
                "gaps — missing GSTIN, TAN, CIN, contact, financial year…"),
    ("", "", ""),
    ("h2", "Fill", ""),
    ("p", "Remaining clients", "Rows that already have a Client code: fill only the blanks. A blank cell never "
                               "wipes what the portal already has."),
    ("p", "New clients", "Add a row with Client code left blank. Minimum: Legal name and either the GSTIN or the "
                         "PAN. Type the GSTIN and the PAN, state and entity type are worked out for you."),
    ("p", "Engagements", "Pick the Financial year and say Yes to each service (GST, TDS / TCS, Income Tax, "
                         "ROC / MCA, Bookkeeping, MIS). The import opens one Compliance engagement per client "
                         "with those services. GST needs a GSTIN and a return frequency; TDS needs a TAN; ROC "
                         "needs a CIN / LLPIN and is only for companies and LLPs."),
    ("p", "Several GSTINs", "Same PAN in several states = ONE client. Put the main GSTIN in 'GSTIN (main)' and "
                            "the rest in 'Other GSTINs', separated by commas."),
    ("p", "Logins", "Optional. Each client's main logins — GST, Income-tax (PAN), Income-tax (TAN), TRACES, "
                    "MCA, E-way bill — go in the 'Logins' columns of its row on 'Clients'. Logins of other "
                    "GSTINs, TRACES taxpayer and E-invoice go on the 'Logins' sheet. Passwords are encrypted "
                    "on import."),
    ("p", "Reading logins", "Export client master masks every password as '(saved in portal)' — importing that "
                            "back changes nothing. Export with passwords writes them in full for the clients "
                            "whose logins you may see (your office as partner, or your engagements); the rest "
                            "stay masked. Every password exported is logged against your name."),
    ("", "", ""),
    ("h2", "Import (Excel → portal)", ""),
    ("p", "1.", "Clients → Import client master, then choose this file."),
    ("p", "2.", "The portal shows a preview first: new clients, updates (only the cells that change), "
                "engagements it will open, and every row it will skip and why. Nothing is saved yet."),
    ("p", "3.", "Confirm. Rows marked ✖ are skipped; fix them and import the same file again — "
                "importing twice never creates duplicates (rows are matched by Client code, then GSTIN, then PAN)."),
    ("", "", ""),
    ("h2", "Rules", ""),
    ("p", "Do not", "Rename, reorder or delete the column headings on 'Clients' and 'Logins', or edit the 'Lists' sheet."),
    ("p", "Dates", "dd-mm-yyyy."),
    ("p", "Entity type guess", "From the PAN's 4th letter: C company, P proprietorship, F firm, H HUF, A/B AOP, "
                               "T trust. Pick the real type in 'Entity type' if the guess is wrong "
                               "(e.g. LLP, public company, individual)."),
    ("p", "Template version", "client-master v2"),
]
for i, (kind, b, c) in enumerate(rows, start=1):
    rm.cell(row=i, column=2, value=b)
    rm.cell(row=i, column=3, value=c)
    cb, cc = rm.cell(row=i, column=2), rm.cell(row=i, column=3)
    cb.font, cc.font = f(bold=True), f()
    cc.alignment = WRAP
    cb.alignment = Alignment(vertical="top")
    if kind == "h1":
        cb.font = f(bold=True, size=16, color="1F3A5F")
        rm.row_dimensions[i].height = 26
    elif kind == "h2":
        cb.font = f(bold=True, size=12, color="FFFFFF")
        for col in (2, 3):
            rm.cell(row=i, column=col).fill = TITLE
    elif kind in ("in", "auto", "ex"):
        cb.fill = {"in": INPUT, "auto": AUTO, "ex": EXAMPLE}[kind]
        cb.border = BORDER
    if c and len(c) > 95:
        rm.row_dimensions[i].height = 14 * (len(c) // 90 + 1)
rm.sheet_view.showGridLines = False
rm.sheet_properties.tabColor = "2E5E4E"

for ws in (rm, cs, lg):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.active = 1  # open on Clients
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT)
print("saved", OUT, "example GSTIN", EX_GSTIN)
