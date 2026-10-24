"""Builds the shareable Client Master v2 file from the peer's merged workbook
and the firm's login workbook. Never prints a password.

    python -I build_v2_file.py TEMPLATE MERGED LOGINS OUT
"""
import re
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Font

TEMPLATE, MERGED, LOGINS, OUT = sys.argv[1:5]
SAVED = "(saved in portal)"

tpl = openpyxl.load_workbook(TEMPLATE)
src = openpyxl.load_workbook(MERGED, data_only=True)
lw = openpyxl.load_workbook(LOGINS, data_only=True)

cs, ss = tpl["Clients"], src["Clients"]
head = {cs.cell(row=4, column=c).value: c for c in range(1, cs.max_column + 1)}
src_head = {ss.cell(row=4, column=c).value: c for c in range(1, ss.max_column + 1)}
V1 = [h for h in list(head)[:31]]
assert V1[0] == "Client code" and V1[30] == "Lead / in-charge (email)", V1

# ── Clients: the merged rows, as they are ────────────────────────────────
rows = []  # (sheet row, dict of values)
out_r = 6
for r in range(6, ss.max_row + 1):
    vals = {h: ss.cell(row=r, column=src_head[h]).value for h in V1 if h in src_head}
    if all(v in (None, "") for v in vals.values()):
        continue
    for h, v in vals.items():
        if v not in (None, ""):
            cs.cell(row=out_r, column=head[h]).value = v
    rows.append((out_r, vals))
    out_r += 1
print("client rows", len(rows), "last row", out_r - 1)

bare = lambda s: re.sub(r"[^0-9A-Z]", "", str(s or "").upper())
namekey = lambda s: re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def pan_of(vals):
    p = bare(vals.get("PAN"))
    if not p and len(bare(vals.get("GSTIN (main)"))) == 15:
        p = bare(vals.get("GSTIN (main)"))[2:12]
    return p


by_main_gstin, by_other_gstin, by_pan, by_tan, by_name = {}, {}, {}, {}, {}
for r, v in rows:
    g = bare(v.get("GSTIN (main)"))
    if g:
        by_main_gstin.setdefault(g, r)
    for o in re.split(r"[\s,;/]+", str(v.get("Other GSTINs") or "")):
        if bare(o):
            by_other_gstin.setdefault(bare(o), r)
    if pan_of(v):
        by_pan.setdefault(pan_of(v), []).append(r)
    if bare(v.get("TAN")):
        by_tan.setdefault(bare(v.get("TAN")), r)
    by_name.setdefault(namekey(v.get("Legal name *")), []).append(r)
vals_of = dict(rows)


def cell(r, h):
    return cs.cell(row=r, column=head[h])


def put_login(r, user_h, pw_h, user, pw):
    """On the row, unless it already holds a real login there. True when placed."""
    have = cell(r, pw_h).value
    if have not in (None, "", SAVED):
        return False
    cell(r, user_h).value = str(user)
    cell(r, pw_h).value = str(pw)
    return True


def by_client_name(name):
    """The one Clients row of this name: exact, then without "(…)", then a unique prefix."""
    for key in (namekey(name), namekey(re.sub(r"\(.*?\)", "", str(name)))):
        if len(by_name.get(key, [])) == 1:
            return by_name[key][0]
    key = namekey(re.sub(r"\(.*?\)", "", str(name)))
    hits = {rr for k, lst in by_name.items() for rr in lst
            if k and len(key) >= 6 and (k.startswith(key) or key.startswith(k))}
    return hits.pop() if len(hits) == 1 else None


extra_logins = []  # (code, name, portal, target, user, password)
notes = []
counts = dict(gst_by_name=0, gst_row=0, gst_sheet=0, it=0, tan=0, traces=0, taxpayer=0, masked_row=0,
              masked_sheet=0, tan_filled=0)
unmatched = []

# ── GST logins: by GSTIN ─────────────────────────────────────────────────
gs = lw["GST Login"]
gh = {str(gs.cell(row=2, column=c).value or "").strip(): c for c in range(1, 8)}
real_gst = set()
for r in range(3, gs.max_row + 1):
    name = gs.cell(row=r, column=gh["Particular"]).value
    user = gs.cell(row=r, column=gh["ID"]).value
    pw = gs.cell(row=r, column=gh["Pwd"]).value
    g = bare(gs.cell(row=r, column=gh["GSTIN"]).value)
    if not (user and pw) or name in (None, "Particular"):
        continue
    if not g:
        # The sheet often leaves the GSTIN out: the client, by its name.
        row = by_client_name(name)
        if row is None:
            unmatched.append(f"GST login of {name} (no GSTIN): no single Clients row has this name")
            continue
        if put_login(row, "GST user ID", "GST password", user, pw):
            counts["gst_row"] += 1
            counts["gst_by_name"] += 1
            if bare(cell(row, "GSTIN (main)").value):
                real_gst.add(bare(cell(row, "GSTIN (main)").value))
        continue
    if g in by_main_gstin:
        if put_login(by_main_gstin[g], "GST user ID", "GST password", user, pw):
            counts["gst_row"] += 1
            real_gst.add(g)
    elif g in by_other_gstin:
        cr = by_other_gstin[g]
        extra_logins.append((cell(cr, "Client code").value, vals_of[cr]["Legal name *"], "GST", g, user, pw))
        counts["gst_sheet"] += 1
        real_gst.add(g)
    else:
        unmatched.append(f"GST login of {name} ({g or 'no GSTIN'}): no Clients row has this GSTIN")

# ── Income-tax / TRACES logins: by PAN, TAN, then name ───────────────────
ts = lw["TDS Login"]
for r in range(2, ts.max_row + 1):
    get = lambda c: ts.cell(row=r, column=c).value
    name = get(2)
    if not name:
        continue
    it_user, it_pw, tan_user, tan_pw, tr_user, tr_pw, tp_user, tp_pw = (get(c) for c in range(4, 12))
    target = None
    p = bare(it_user)
    if re.fullmatch(r"[A-Z]{5}[0-9]{4}[A-Z]", p) and len(by_pan.get(p, [])) == 1:
        target = by_pan[p][0]
    if target is None and bare(tan_user) in by_tan:
        target = by_tan[bare(tan_user)]
    if target is None:
        target = by_client_name(name)
    if target is None:
        unmatched.append(f"Income-tax / TRACES logins of {name}: no single Clients row matches")
        continue
    if it_user and it_pw and put_login(target, "Income-tax user ID (PAN)", "Income-tax password (PAN)", it_user, it_pw):
        counts["it"] += 1
    if tan_user and tan_pw:
        t = bare(tan_user)
        if not bare(cell(target, "TAN").value) and re.fullmatch(r"[A-Z]{4}[0-9]{5}[A-Z]", t):
            cell(target, "TAN").value = t
            counts["tan_filled"] += 1
            notes.append(f"{vals_of[target]['Legal name *']}: TAN {t} filled in from its income-tax login.")
        if put_login(target, "Income-tax user ID (TAN)", "Income-tax password (TAN)", tan_user, tan_pw):
            counts["tan"] += 1
    if tr_user and tr_pw and put_login(target, "TRACES user ID (TAN)", "TRACES password", tr_user, tr_pw):
        counts["traces"] += 1
    if tp_user and tp_pw:
        extra_logins.append((cell(target, "Client code").value, vals_of[target]["Legal name *"],
                             "TRACES taxpayer", "", tp_user, tp_pw))
        counts["taxpayer"] += 1

# ── The merged file's masked logins (already in the portal) ─────────────
ROW_COLUMNS = {  # Logins-sheet portal → (user column, password column, the number it is for)
    "GST": ("GST user ID", "GST password", "GSTIN (main)"),
    "Income-tax (PAN)": ("Income-tax user ID (PAN)", "Income-tax password (PAN)", None),
    "Income-tax (TAN)": ("Income-tax user ID (TAN)", "Income-tax password (TAN)", "TAN"),
    "TRACES deductor": ("TRACES user ID (TAN)", "TRACES password", "TAN"),
    "MCA": ("MCA user ID", "MCA password", None),
    "E-way bill": ("E-way bill user ID", "E-way bill password", "GSTIN (main)"),
}
code_row = {str(cell(r, "Client code").value).upper(): r for r, _ in rows if cell(r, "Client code").value}
sl = src["Logins"]
for r in range(5, sl.max_row + 1):
    code, name, portal, target, user, pw = (sl.cell(row=r, column=c).value for c in range(1, 7))
    if not (user and pw) or str(code).upper() == "EXAMPLE":
        continue
    g = bare(target)
    row = code_row.get(str(code or "").upper())
    col = ROW_COLUMNS.get(portal)
    if row and col:
        user_h, pw_h, on = col
        main = bare(cell(row, on).value) if on else ""
        if not g or not on or g == main or not main:
            # The client's main login: on its row, unless the real one is there already.
            if cell(row, pw_h).value in (None, ""):
                cell(row, user_h).value, cell(row, pw_h).value = user, pw
                counts["masked_row"] += 1
            continue
    if any(e[0] == code and e[2] == portal and bare(e[3]) == g for e in extra_logins):
        continue  # the real one is on the Logins sheet
    extra_logins.append((code, name, portal, target or "", user, pw))
    counts["masked_sheet"] += 1

lg = tpl["Logins"]
for i, (code, name, portal, target, user, pw) in enumerate(extra_logins):
    for c, v in enumerate((code, name, portal, target, user, pw), start=1):
        if v not in (None, ""):
            lg.cell(row=5 + i, column=c).value = str(v)

# ── The peer's sheets, carried over ──────────────────────────────────────
for name in ("Review before import", "Merge notes"):
    s = src[name]
    d = tpl.create_sheet(name, index=tpl.sheetnames.index("Lists"))
    for row in s.iter_rows():
        for c in row:
            n = d.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font, n.fill, n.alignment, n.border = copy(c.font), copy(c.fill), copy(c.alignment), copy(c.border)
                n.number_format = c.number_format
    for k, dim in s.column_dimensions.items():
        d.column_dimensions[k].width = dim.width
    for rng in s.merged_cells.ranges:
        d.merge_cells(str(rng))
    d.sheet_properties.tabColor = s.sheet_properties.tabColor

mn = tpl["Merge notes"]
r = mn.max_row + 2
lines = [
    ("Logins (client-master v2)", None),
    ("Format", "Each client's main logins are on its own row on 'Clients' (columns AF:AQ). Logins of "
               "other GSTINs and TRACES taxpayer are on the 'Logins' sheet."),
    ("Source", "Client GST Login Details - DHVAJ clean.xlsx (sheets 'GST Login' and 'TDS Login')."),
    ("GST logins on the client's row", f"{counts['gst_row']} ({counts['gst_by_name']} matched by name: "
                                       "the login sheet has no GSTIN for them)"),
    ("GST logins of another GSTIN (Logins sheet)", str(counts["gst_sheet"])),
    ("Income-tax (PAN) logins", str(counts["it"])),
    ("Income-tax (TAN) logins", str(counts["tan"])),
    ("TRACES deductor logins", str(counts["traces"])),
    ("TRACES taxpayer logins (Logins sheet)", str(counts["taxpayer"])),
    ("TANs filled in from a login", str(counts["tan_filled"])),
    ("Already in the portal, kept masked", f"{counts['masked_row'] + counts['masked_sheet']} "
                                           "('(saved in portal)' — importing them changes nothing)"),
    ("Not placed", str(len(unmatched)) + (" — " + "; ".join(unmatched) if unmatched else "")),
    ("Keep private", "This file holds real passwords. Share it only with people who may see them, "
                     "import it, then delete it. The portal encrypts each password on import."),
] + [("TAN note", n) for n in notes]
for k, v in lines:
    a = mn.cell(row=r, column=1, value=k)
    a.font = Font(name="Arial", bold=True, size=12 if v is None else 10)
    if v is not None:
        b = mn.cell(row=r, column=2, value=v)
        b.font = Font(name="Arial", size=10)
    r += 1

tpl.active = tpl.sheetnames.index("Clients")
tpl.save(OUT)
print({k: v for k, v in counts.items()}, "extra logins", len(extra_logins), "unmatched", len(unmatched))
for u in unmatched:
    print("  unmatched:", u)
